"""Phase 5 tests: a generated document whose numbers cannot be traced is a bug."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest

from campaign_copilot.reporting import (
    Figure,
    FigureSet,
    ProvenanceError,
    build_weekly_review,
    render_markdown,
    render_pptx,
)
from campaign_copilot.reporting.render import ROWS_PER_APPENDIX_SLIDE
from campaign_copilot.semantic.layer import SemanticLayer
from campaign_copilot.tools.sql import Warehouse

DB = Path(__file__).resolve().parents[1] / "warehouse" / "campaign_copilot.duckdb"
WEEK_END = date(2025, 12, 31)

pytestmark = pytest.mark.skipif(not DB.exists(), reason="run `make warehouse`")


@pytest.fixture(scope="module")
def warehouse() -> Warehouse:
    return Warehouse(DB)


@pytest.fixture(scope="module")
def review(warehouse: Warehouse):
    return build_weekly_review(SemanticLayer.load(), warehouse, WEEK_END)


# ------------------------------------------------------------------- figures


def test_a_figure_without_provenance_cannot_exist() -> None:
    """A number with no query and no parents is a claim, which is the thing we avoid."""
    with pytest.raises(ProvenanceError, match="neither SQL nor parents"):
        Figure(id="X", label="made up", value=42.0)


def test_a_derived_figure_needs_a_formula() -> None:
    with pytest.raises(ProvenanceError, match="no formula"):
        Figure(id="X", label="x", value=1.0, derived_from=("A",))


def test_a_derived_figure_cannot_cite_an_unregistered_parent() -> None:
    figures = FigureSet()
    with pytest.raises(ProvenanceError, match="not registered"):
        figures.add(
            Figure(id="B", label="b", value=1.0, derived_from=("A",), formula="pct_change")
        )


def test_duplicate_figure_ids_are_rejected() -> None:
    figures = FigureSet()
    figures.add(Figure(id="F1", label="a", value=1.0, sql="select 1"))
    with pytest.raises(ProvenanceError, match="Duplicate"):
        figures.add(Figure(id="F1", label="b", value=2.0, sql="select 2"))


def test_derived_arithmetic_is_a_whitelist_not_eval() -> None:
    """A report generator that evaluates expressions from a data file has an RCE bug.

    Rejected at construction, not at verification: a Figure that cannot be trusted should
    never exist, rather than exist and fail a later check.
    """
    with pytest.raises(ProvenanceError, match="whitelist, not an expression"):
        Figure(
            id="C",
            label="c",
            value=99.0,
            derived_from=("A", "B"),
            formula="__import__('os').system('id')",
        )


def test_verify_catches_a_figure_that_no_longer_reproduces() -> None:
    figures = FigureSet()
    figures.add(Figure(id="F1", label="spend", value=100.0, sql="select 999"))
    problems = figures.verify(lambda sql: [[float(sql.split()[-1])]])
    assert problems and "SQL yields" in problems[0]


def test_verify_catches_a_derived_figure_that_drifted() -> None:
    figures = FigureSet()
    figures.add(Figure(id="A", label="a", value=110.0, sql="select 110"))
    figures.add(Figure(id="B", label="b", value=100.0, sql="select 100"))
    figures.add(
        Figure(id="C", label="c", value=50.0, derived_from=("A", "B"), formula="pct_change")
    )
    problems = figures.verify(lambda sql: [[float(sql.split()[-1])]])
    assert any(p.startswith("C:") for p in problems)


# ------------------------------------------------- the report against real data


def test_every_figure_reproduces_from_its_own_provenance(review, warehouse: Warehouse) -> None:
    """Re-run every query and recompute every derivation. Nothing may have drifted."""
    assert review.figures.verify(lambda sql: warehouse.execute(sql)[1]) == []


def test_the_narrative_states_only_numbers_the_figures_support(review) -> None:
    from campaign_copilot.grounding import GroundingChecker

    assert GroundingChecker().check(review.narrative, review.figures.facts()).ok


def test_the_narrative_footnotes_every_number(review) -> None:
    for fid in ("F1", "F2", "F3", "F4", "F7", "F8"):
        assert f"[{fid}]" in review.narrative


def test_an_ungrounded_narrative_fails_the_build(monkeypatch: pytest.MonkeyPatch) -> None:
    """The last place a wrong number could reach a client. It does not get to.

    The agent refuses to say an ungrounded number; the report refuses to print one; and both
    refusals go through the same GroundingChecker, not two implementations of one idea.
    """
    from campaign_copilot.reporting import builder

    monkeypatch.setattr(
        builder, "_narrative", lambda review, best: "Spend was $412,000.00 this week."
    )
    with pytest.raises(ProvenanceError, match=r"\$412,000"):
        build_weekly_review(SemanticLayer.load(), Warehouse(DB), WEEK_END)


def test_direct_traffic_is_excluded_from_every_headline_figure(review) -> None:
    for fid in ("F1", "F2", "F3", "F4"):
        sql = review.figures[fid].sql
        assert "(direct)" in sql, f"{fid} does not exclude unattributed traffic"


def test_the_exclusion_is_stated_in_the_prose_not_just_the_sql(review) -> None:
    assert "direct" in review.narrative.lower()
    assert "excluded" in review.narrative.lower()


# ----------------------------------------------------------------- rendering


def test_markdown_carries_the_full_appendix(review) -> None:
    markdown = render_markdown(review)
    for fid in review.figures.figures:
        assert f"`{fid}`" in markdown
    assert "## Provenance" in markdown
    assert "sum(revenue_usd)" in markdown, "the appendix must show real SQL, not a summary"


def test_markdown_lists_every_channel(review) -> None:
    markdown = render_markdown(review)
    for row in review.channels:
        assert row.channel in markdown


def test_pptx_renders_and_paginates_the_appendix(review, tmp_path: Path) -> None:
    from pptx import Presentation

    path = render_pptx(review, tmp_path / "deck.pptx")
    prs = Presentation(str(path))
    pages = -(-len(review.figures.figures) // ROWS_PER_APPENDIX_SLIDE)
    assert len(prs.slides) == 3 + pages, "title, headline, chart, then appendix pages"


def test_the_deck_puts_full_sql_in_the_speaker_notes(review, tmp_path: Path) -> None:
    """The slide is not the archive: SQL is truncated on it, and complete underneath."""
    from pptx import Presentation

    path = render_pptx(review, tmp_path / "deck.pptx")
    prs = Presentation(str(path))
    notes = "\n".join(
        s.notes_slide.notes_text_frame.text for s in prs.slides if s.has_notes_slide
    )
    assert "nullif(sum(spend_usd), 0)" in notes


def test_the_deck_contains_no_placeholder_text(review, tmp_path: Path) -> None:
    from pptx import Presentation

    path = render_pptx(review, tmp_path / "deck.pptx")
    prs = Presentation(str(path))
    text = " ".join(
        shape.text_frame.text
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame
    ).lower()
    for banned in ("lorem", "ipsum", "todo", "[insert", "xxx"):
        assert banned not in text
